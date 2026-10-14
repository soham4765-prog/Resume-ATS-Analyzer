
from io import BytesIO
from docx import Document


def extract_docx_text(file_bytes: bytes) -> str:
    """Extract paragraphs and table text from a DOCX file."""

    if not file_bytes:
        raise ValueError("The uploaded DOCX file is empty.")

    try:
        document = Document(BytesIO(file_bytes))

        parts = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        # Resume details may be stored in tables.
        for table in document.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip()
                    for cell in row.cells
                ]

                parts.append(" | ".join(cells))

        extracted_text = "\n".join(parts).strip()

    except Exception as exc:
        raise ValueError(
            "The DOCX file could not be read."
        ) from exc

    if not extracted_text:
        raise ValueError(
            "No readable text was found in the DOCX file."
        )

    return extracted_text
